"""月度「已收票 + 对公付款」清单。

以已经落库的银行付款↔进项发票关联为唯一付款事实，再向采购订单与商品分配明细下钻。
不复制付款匹配逻辑；未匹配的银行流水/发票不会进入本清单。
"""
from __future__ import annotations

from collections import defaultdict
from datetime import date
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.models.bank import BankAccount, BankTransaction
from app.models.purchase import ExternalPurchaseOrder, PurchaseAllocationItem, Supplier
from app.models.tax import TaxInvoice, TaxInvoiceLink
from app.services import invoice_reconciliation

TOLERANCE = Decimal("0.05")
TARGET_TYPE = "bank_transaction"


def _dec(value: Decimal | float | int | str | None) -> Decimal:
    if value is None:
        return Decimal("0")
    if isinstance(value, Decimal):
        return value
    return Decimal(str(value))


def _month_range(year: int, month: int) -> tuple[date, date]:
    if month < 1 or month > 12:
        raise ValueError("月份必须在 1-12")
    start = date(year, month, 1)
    end = date(year + (month == 12), 1 if month == 12 else month + 1, 1)
    return start, end


def _active_bank_links(db: Session, *, invoice_ids: list[int] | None = None) -> list[TaxInvoiceLink]:
    q = (
        db.query(TaxInvoiceLink)
        .filter(
            TaxInvoiceLink.target_type == TARGET_TYPE,
            TaxInvoiceLink.confirmed.is_(True),
            TaxInvoiceLink.match_method != "rejected",
        )
    )
    if invoice_ids:
        q = q.filter(TaxInvoiceLink.invoice_id.in_(invoice_ids))
    return q.all()


def _company_matches(invoice: TaxInvoice, company: str) -> bool:
    """买方名称有值时按公司主体隔离；历史票缺买方名称时不误删。"""
    if not company or not (invoice.buyer_name or "").strip():
        return True
    return (
        invoice_reconciliation.normalize_supplier(invoice.buyer_name)
        == invoice_reconciliation.normalize_supplier(company)
    )


def _invoice_purchase_map(db: Session, invoices: list[TaxInvoice]) -> dict[int, list[dict[str, Any]]]:
    """复用采购发票对账的人工优先 + 截止开票日 FIFO 结果。"""
    by_supplier: dict[str, list[TaxInvoice]] = defaultdict(list)
    for inv in invoices:
        key = invoice_reconciliation.normalize_supplier(inv.seller_name)
        if key:
            by_supplier[key].append(inv)

    result: dict[int, list[dict[str, Any]]] = {}
    for rows in by_supplier.values():
        seller = rows[0].seller_name
        reconciliation = invoice_reconciliation.reconcile(db, supplier=seller)
        for supplier_block in reconciliation.get("suppliers", []):
            for month_block in supplier_block.get("months", []):
                for inv_row in month_block.get("invoices", []):
                    result[int(inv_row["invoiceId"])] = list(inv_row.get("covered") or [])
    return result


def build_report(db: Session, year: int, month: int, company: str = "") -> dict[str, Any]:
    start, end = _month_range(year, month)

    txns = (
        db.query(BankTransaction)
        .filter(
            BankTransaction.txn_date >= start,
            BankTransaction.txn_date < end,
            BankTransaction.direction == "out",
        )
        .order_by(BankTransaction.txn_date, BankTransaction.id)
        .all()
    )
    txn_map = {row.id: row for row in txns}
    if not txn_map:
        return {
            "year": year,
            "month": month,
            "company": company,
            "summary": {
                "paymentCount": 0,
                "invoiceCount": 0,
                "paymentTotal": "0.00",
                "allocatedTotal": "0.00",
                "invoiceTotal": "0.00",
                "outstandingTotal": "0.00",
                "productRowCount": 0,
            },
            "rows": [],
            "productDetails": [],
        }

    links = (
        db.query(TaxInvoiceLink)
        .filter(
            TaxInvoiceLink.target_type == TARGET_TYPE,
            TaxInvoiceLink.target_id.in_(txn_map),
            TaxInvoiceLink.confirmed.is_(True),
            TaxInvoiceLink.match_method != "rejected",
        )
        .order_by(TaxInvoiceLink.id)
        .all()
    )
    invoice_ids = sorted({row.invoice_id for row in links})
    invoices = (
        db.query(TaxInvoice)
        .filter(TaxInvoice.id.in_(invoice_ids), TaxInvoice.direction == "input")
        .all()
        if invoice_ids
        else []
    )
    invoice_map = {
        inv.id: inv
        for inv in invoices
        if _company_matches(inv, company)
    }
    links = [link for link in links if link.invoice_id in invoice_map]
    invoice_ids = sorted(invoice_map)

    account_ids = sorted({txn.account_id for txn in txns if txn.account_id})
    accounts = (
        db.query(BankAccount).filter(BankAccount.id.in_(account_ids)).all()
        if account_ids
        else []
    )
    account_map = {row.id: row for row in accounts}

    all_invoice_links = _active_bank_links(db, invoice_ids=invoice_ids) if invoice_ids else []
    invoice_bank_total: dict[int, Decimal] = defaultdict(Decimal)
    txn_bank_total: dict[int, Decimal] = defaultdict(Decimal)
    for link in all_invoice_links:
        amount = _dec(link.allocated_amount)
        invoice_bank_total[link.invoice_id] += amount
        txn_bank_total[link.target_id] += amount

    supplier_tax: dict[str, str] = {}
    for supplier in db.query(Supplier).all():
        key = invoice_reconciliation.normalize_supplier(supplier.name)
        if key and supplier.tax_no:
            supplier_tax.setdefault(key, supplier.tax_no)

    purchase_map = _invoice_purchase_map(db, list(invoice_map.values())) if invoice_map else {}
    order_ids = sorted({
        int(order["orderId"])
        for covered in purchase_map.values()
        for order in covered
        if order.get("orderId") is not None
    })
    orders = (
        db.query(ExternalPurchaseOrder).filter(ExternalPurchaseOrder.id.in_(order_ids)).all()
        if order_ids
        else []
    )
    order_map = {row.id: row for row in orders}
    items = (
        db.query(PurchaseAllocationItem)
        .filter(PurchaseAllocationItem.po_id.in_(order_ids))
        .order_by(PurchaseAllocationItem.po_id, PurchaseAllocationItem.id)
        .all()
        if order_ids
        else []
    )
    items_by_order: dict[int, list[PurchaseAllocationItem]] = defaultdict(list)
    for item in items:
        items_by_order[item.po_id].append(item)

    rows: list[dict[str, Any]] = []
    for link in links:
        txn = txn_map.get(link.target_id)
        inv = invoice_map.get(link.invoice_id)
        if txn is None or inv is None:
            continue
        account = account_map.get(txn.account_id)
        allocated = _dec(link.allocated_amount)
        inv_total = _dec(inv.total_amount)
        paid_total = invoice_bank_total.get(inv.id, Decimal("0"))
        txn_total = _dec(txn.amount)
        txn_allocated = txn_bank_total.get(txn.id, Decimal("0"))
        covered = purchase_map.get(inv.id, [])
        order_nos = [str(x.get("orderNo") or "") for x in covered if x.get("orderNo")]
        rows.append({
            "linkId": link.id,
            "paymentId": txn.id,
            "paymentDate": txn.txn_date.isoformat(),
            "paymentAccount": account.account_no if account else "",
            "paymentAccountName": account.account_name if account else "",
            "supplierName": inv.seller_name or txn.counterparty_name,
            "supplierTaxId": inv.seller_tax_id or supplier_tax.get(
                invoice_reconciliation.normalize_supplier(inv.seller_name), ""
            ),
            "counterpartyAccount": txn.counterparty_account,
            "voucherNo": txn.voucher_no,
            "summary": txn.summary,
            "paymentAmount": str(txn_total),
            "paymentAllocatedAmount": str(allocated),
            "paymentMatchedTotal": str(txn_allocated),
            "paymentStatus": "matched" if txn_total - txn_allocated <= TOLERANCE else "partial",
            "invoiceId": inv.id,
            "invoiceNumber": inv.invoice_number,
            "invoiceDate": inv.issue_date.strftime("%Y-%m-%d") if inv.issue_date else "",
            "invoiceType": inv.invoice_type,
            "invoiceAmountExclTax": str(_dec(inv.amount_excl_tax)),
            "invoiceTaxAmount": str(_dec(inv.tax_amount)),
            "invoiceTotalAmount": str(inv_total),
            "invoiceCorporatePaidTotal": str(paid_total),
            "invoiceOutstandingAmount": str(max(inv_total - paid_total, Decimal("0"))),
            "invoiceStatus": "paid" if inv_total - paid_total <= TOLERANCE else "partial",
            "purchaseOrderNos": order_nos,
        })

    product_details: list[dict[str, Any]] = []
    for inv_id, covered in purchase_map.items():
        inv = invoice_map.get(inv_id)
        if inv is None:
            continue
        for coverage in covered:
            po_id = int(coverage.get("orderId") or 0)
            order = order_map.get(po_id)
            order_no = str(coverage.get("orderNo") or (order.external_order_id if order else ""))
            covered_amount = _dec(coverage.get("consumed"))
            item_rows = items_by_order.get(po_id) or [None]
            for item in item_rows:
                product_details.append({
                    "invoiceId": inv.id,
                    "invoiceNumber": inv.invoice_number,
                    "invoiceDate": inv.issue_date.strftime("%Y-%m-%d") if inv.issue_date else "",
                    "supplierName": inv.seller_name,
                    "purchaseOrderId": po_id or None,
                    "purchaseOrderNo": order_no,
                    "platform": (order.platform if order else str(coverage.get("platform") or "")),
                    "invoiceCoveredOrderAmount": str(covered_amount),
                    "skuCode": item.sku_code if item else "",
                    "productName": item.goods_name if item else (order.title if order else ""),
                    "quantity": str(_dec(item.quantity)) if item else "",
                    "unitPrice": str(_dec(item.unit_price)) if item else "",
                    "itemAmount": str(_dec(item.amount)) if item else "",
                })

    unique_txns = {row["paymentId"] for row in rows}
    unique_invoices = {row["invoiceId"] for row in rows}
    payment_total = sum((_dec(txn_map[i].amount) for i in unique_txns), Decimal("0"))
    allocated_total = sum((_dec(row["paymentAllocatedAmount"]) for row in rows), Decimal("0"))
    invoice_total = sum((_dec(invoice_map[i].total_amount) for i in unique_invoices), Decimal("0"))
    outstanding_total = sum(
        (max(_dec(invoice_map[i].total_amount) - invoice_bank_total.get(i, Decimal("0")), Decimal("0"))
         for i in unique_invoices),
        Decimal("0"),
    )

    return {
        "year": year,
        "month": month,
        "company": company,
        "summary": {
            "paymentCount": len(unique_txns),
            "invoiceCount": len(unique_invoices),
            "paymentTotal": str(payment_total),
            "allocatedTotal": str(allocated_total),
            "invoiceTotal": str(invoice_total),
            "outstandingTotal": str(outstanding_total),
            "productRowCount": len(product_details),
        },
        "rows": rows,
        "productDetails": product_details,
    }


def _style_sheet(ws) -> None:
    ws.freeze_panes = "A2"
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        width = max(
            [len(str(ws.cell(row=r, column=col).value or "")) for r in range(1, min(ws.max_row, 80) + 1)]
            or [10]
        )
        ws.column_dimensions[letter].width = min(max(width + 2, 10), 34)


def corporate_payment_xlsx(report: dict[str, Any]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "对公付款汇总"
    headers = [
        "付款日期", "供应商/对方", "供应商税号", "我方付款账号", "我方账户名", "对方账号",
        "银行流水/凭证号", "付款摘要", "付款金额", "已匹配发票金额", "未对账差额",
        "付款对账状态", "已匹配发票号码", "发票张数", "关联采购订单",
    ]
    ws.append(headers)

    payments: dict[int, dict[str, Any]] = {}
    for row in report.get("rows", []):
        payment_id = int(row["paymentId"])
        group = payments.setdefault(payment_id, {
            "paymentDate": row["paymentDate"],
            "supplierNames": [],
            "supplierTaxIds": [],
            "paymentAccount": row["paymentAccount"],
            "paymentAccountName": row["paymentAccountName"],
            "counterpartyAccount": row["counterpartyAccount"],
            "voucherNo": row["voucherNo"],
            "summary": row["summary"],
            "paymentAmount": _dec(row["paymentAmount"]),
            "paymentMatchedTotal": _dec(row["paymentMatchedTotal"]),
            "paymentStatus": row["paymentStatus"],
            "invoiceNumbers": [],
            "purchaseOrderNos": [],
        })
        if row["supplierName"] and row["supplierName"] not in group["supplierNames"]:
            group["supplierNames"].append(row["supplierName"])
        if row["supplierTaxId"] and row["supplierTaxId"] not in group["supplierTaxIds"]:
            group["supplierTaxIds"].append(row["supplierTaxId"])
        if row["invoiceNumber"] and row["invoiceNumber"] not in group["invoiceNumbers"]:
            group["invoiceNumbers"].append(row["invoiceNumber"])
        for order_no in row["purchaseOrderNos"]:
            if order_no and order_no not in group["purchaseOrderNos"]:
                group["purchaseOrderNos"].append(order_no)

    for group in payments.values():
        remaining = max(group["paymentAmount"] - group["paymentMatchedTotal"], Decimal("0"))
        ws.append([
            group["paymentDate"],
            "、".join(group["supplierNames"]),
            "、".join(group["supplierTaxIds"]),
            group["paymentAccount"],
            group["paymentAccountName"],
            group["counterpartyAccount"],
            group["voucherNo"],
            group["summary"],
            float(group["paymentAmount"]),
            float(group["paymentMatchedTotal"]),
            float(remaining),
            "已对清" if group["paymentStatus"] == "matched" else "部分对账",
            "、".join(group["invoiceNumbers"]),
            len(group["invoiceNumbers"]),
            "、".join(group["purchaseOrderNos"]),
        ])
    _style_sheet(ws)

    detail = wb.create_sheet("商品明细")
    detail.append([
        "发票号码", "发票日期", "供应商", "采购订单号", "采购平台", "发票覆盖订单金额",
        "商品编码", "商品名称", "数量", "采购单价", "商品金额",
    ])
    for row in report.get("productDetails", []):
        detail.append([
            row["invoiceNumber"], row["invoiceDate"], row["supplierName"],
            row["purchaseOrderNo"], row["platform"], float(_dec(row["invoiceCoveredOrderAmount"])),
            row["skuCode"], row["productName"],
            float(_dec(row["quantity"])) if row["quantity"] != "" else None,
            float(_dec(row["unitPrice"])) if row["unitPrice"] != "" else None,
            float(_dec(row["itemAmount"])) if row["itemAmount"] != "" else None,
        ])
    _style_sheet(detail)

    summary = wb.create_sheet("月度汇总", 0)
    s = report.get("summary", {})
    summary.append(["指标", "数值"])
    summary.append(["对公付款笔数", s.get("paymentCount", 0)])
    summary.append(["进项发票张数", s.get("invoiceCount", 0)])
    summary.append(["对公付款总额", float(_dec(s.get("paymentTotal")))])
    summary.append(["已关联发票金额", float(_dec(s.get("allocatedTotal")))])
    summary.append(["发票价税合计", float(_dec(s.get("invoiceTotal")))])
    summary.append(["发票未付余额", float(_dec(s.get("outstandingTotal")))])
    summary.append(["商品明细行数", s.get("productRowCount", 0)])
    _style_sheet(summary)

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
