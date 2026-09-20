"""月度已收票对公付款清单。"""
from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO
from uuid import uuid4

from openpyxl import load_workbook

from app.models.bank import BankAccount, BankTransaction
from app.models.purchase import ExternalPurchaseOrder, PurchaseAllocationItem, Supplier
from app.models.tax import TaxInvoice, TaxInvoiceLink
from app.services import finance_corporate_payment_report_service as service


def test_build_report_links_bank_invoice_purchase_and_product(db_session):
    token = uuid4().hex[:10]
    seller = f"测试供应商-{token}"
    company = f"测试公司-{token}"

    account = BankAccount(
        account_no=f"ACCT-{token}",
        account_name=company,
        bank_name="测试银行",
        currency="CNY",
    )
    db_session.add(account)
    db_session.flush()

    supplier = Supplier(
        platform="other",
        external_shop_id=f"SUP-{token}",
        name=seller,
        tax_no=f"TAX-{token}",
    )
    db_session.add(supplier)

    po = ExternalPurchaseOrder(
        external_order_id=f"PO-{token}",
        platform="1688",
        supplier_name=seller,
        ordered_at=datetime(2026, 8, 2, tzinfo=timezone.utc),
        order_amount=Decimal("1000.00"),
        paid_amount=Decimal("1000.00"),
        currency="CNY",
        order_status="paid",
        pay_status="paid",
        raw={},
    )
    db_session.add(po)
    db_session.flush()

    item = PurchaseAllocationItem(
        po_id=po.id,
        sku_code=f"SKU-{token}",
        goods_name="测试商品",
        quantity=Decimal("10"),
        unit_price=Decimal("100"),
        amount=Decimal("1000"),
        source="manual",
    )
    db_session.add(item)

    txn = BankTransaction(
        account_id=account.id,
        txn_date=date(2026, 8, 15),
        direction="out",
        amount=Decimal("1000.00"),
        counterparty_name=seller,
        counterparty_account=f"CP-{token}",
        summary="采购货款",
        voucher_no=f"V-{token}",
        fingerprint=f"pytest-corp-pay-{token}",
        raw={},
    )
    db_session.add(txn)

    invoice = TaxInvoice(
        invoice_key=f"pytest-corp-pay-{token}",
        invoice_number=f"INV-{token}",
        direction="input",
        status="issued",
        issue_date=datetime(2026, 8, 10, tzinfo=timezone.utc),
        seller_name=seller,
        seller_tax_id=f"TAX-{token}",
        buyer_name=company,
        amount_excl_tax=Decimal("884.96"),
        tax_amount=Decimal("115.04"),
        total_amount=Decimal("1000.00"),
        source_system="tax_export",
        raw={},
    )
    db_session.add(invoice)
    db_session.flush()

    db_session.add(TaxInvoiceLink(
        invoice_id=invoice.id,
        target_type="bank_transaction",
        target_id=txn.id,
        allocated_amount=Decimal("1000.00"),
        match_method="manual",
        confirmed=True,
        note="pytest",
    ))
    db_session.commit()

    report = service.build_report(db_session, 2026, 8, company=company)

    assert report["summary"]["paymentCount"] == 1
    assert report["summary"]["invoiceCount"] == 1
    assert Decimal(report["summary"]["allocatedTotal"]) == Decimal("1000.00")
    assert report["rows"][0]["invoiceNumber"] == invoice.invoice_number
    assert report["rows"][0]["voucherNo"] == txn.voucher_no
    assert po.external_order_id in report["rows"][0]["purchaseOrderNos"]
    assert report["productDetails"][0]["skuCode"] == item.sku_code
    assert Decimal(report["productDetails"][0]["unitPrice"]) == Decimal("100")

    blob = service.corporate_payment_xlsx(report)
    wb = load_workbook(BytesIO(blob), read_only=True)
    assert wb.sheetnames == ["月度汇总", "对公付款汇总", "商品明细"]
    assert wb["对公付款汇总"]["M2"].value == invoice.invoice_number
    assert wb["商品明细"]["G2"].value == item.sku_code


def test_report_excludes_unconfirmed_or_other_month(db_session):
    token = uuid4().hex[:10]
    seller = f"未确认供应商-{token}"

    txn = BankTransaction(
        txn_date=date(2026, 8, 5),
        direction="out",
        amount=Decimal("300.00"),
        counterparty_name=seller,
        fingerprint=f"pytest-corp-unconfirmed-{token}",
        raw={},
    )
    db_session.add(txn)
    invoice = TaxInvoice(
        invoice_key=f"pytest-corp-unconfirmed-{token}",
        invoice_number=f"INV-U-{token}",
        direction="input",
        status="issued",
        issue_date=datetime(2026, 8, 3, tzinfo=timezone.utc),
        seller_name=seller,
        total_amount=Decimal("300.00"),
        source_system="tax_export",
        raw={},
    )
    db_session.add(invoice)
    db_session.flush()
    db_session.add(TaxInvoiceLink(
        invoice_id=invoice.id,
        target_type="bank_transaction",
        target_id=txn.id,
        allocated_amount=Decimal("300.00"),
        match_method="manual",
        confirmed=False,
    ))
    db_session.commit()

    report = service.build_report(db_session, 2026, 8)
    ids = {row["invoiceId"] for row in report["rows"]}
    assert invoice.id not in ids



def test_corporate_payment_xlsx_groups_links_by_payment():
    report = {
        "summary": {
            "paymentCount": 1,
            "invoiceCount": 2,
            "paymentTotal": "1000.00",
            "allocatedTotal": "1000.00",
            "invoiceTotal": "1000.00",
            "outstandingTotal": "0.00",
            "productRowCount": 0,
        },
        "rows": [
            {
                "paymentId": 10,
                "paymentDate": "2026-08-15",
                "supplierName": "供应商甲",
                "supplierTaxId": "TAX-A",
                "paymentAccount": "ZJRC-001",
                "paymentAccountName": "测试账户",
                "counterpartyAccount": "CP-001",
                "voucherNo": "V-001",
                "summary": "货款",
                "paymentAmount": "1000.00",
                "paymentAllocatedAmount": "400.00",
                "paymentMatchedTotal": "1000.00",
                "paymentStatus": "matched",
                "invoiceNumber": "INV-A",
                "purchaseOrderNos": ["PO-A"],
            },
            {
                "paymentId": 10,
                "paymentDate": "2026-08-15",
                "supplierName": "供应商甲",
                "supplierTaxId": "TAX-A",
                "paymentAccount": "ZJRC-001",
                "paymentAccountName": "测试账户",
                "counterpartyAccount": "CP-001",
                "voucherNo": "V-001",
                "summary": "货款",
                "paymentAmount": "1000.00",
                "paymentAllocatedAmount": "600.00",
                "paymentMatchedTotal": "1000.00",
                "paymentStatus": "matched",
                "invoiceNumber": "INV-B",
                "purchaseOrderNos": ["PO-B"],
            },
        ],
        "productDetails": [],
    }

    wb = load_workbook(BytesIO(service.corporate_payment_xlsx(report)), read_only=True)
    ws = wb["对公付款汇总"]

    assert ws.max_row == 2
    assert ws["I2"].value == 1000
    assert ws["J2"].value == 1000
    assert ws["K2"].value == 0
    assert ws["M2"].value == "INV-A、INV-B"
    assert ws["N2"].value == 2
    assert ws["O2"].value == "PO-A、PO-B"
